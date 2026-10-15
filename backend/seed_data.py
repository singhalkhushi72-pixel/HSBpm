import os
import sys
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent.parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

import openpyxl
from backend.database import get_db_connection, init_db

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

SEED_FILE = BASE_DIR / "data" / "events_seed.xlsx"

def seed_events():
    """Seeds the events table from data/events_seed.xlsx if not already seeded."""
    init_db()

    if not SEED_FILE.exists():
        raise FileNotFoundError(
            f"Seed data file '{SEED_FILE}' is missing! "
            "Please copy the provided events_seed.xlsx file into the data/ folder before running this script."
        )

    conn = get_db_connection()
    cursor = conn.cursor()

    # Skip seeding if the table already has rows
    cursor.execute("SELECT COUNT(*) AS count FROM events")
    row = cursor.fetchone()
    if row and row["count"] > 0:
        print(f"Events table already has {row['count']} rows. Skipping seed.")
        conn.close()
        return

    # Load workbook and Events sheet
    wb = openpyxl.load_workbook(SEED_FILE)
    if "Events" not in wb.sheetnames:
        raise ValueError(f"Sheet 'Events' not found in {SEED_FILE}. Available sheets: {wb.sheetnames}")

    sheet = wb["Events"]
    rows = list(sheet.iter_rows(values_only=True))

    if not rows:
        raise ValueError("The 'Events' sheet in events_seed.xlsx is empty.")

    header = [str(col).strip().lower() if col is not None else "" for col in rows[0]]
    expected_header = ["id", "title", "category", "description", "event_date", "event_time", "venue", "capacity", "fee", "icon"]

    # Check header
    if header != expected_header:
        print(f"Warning: Header row {header} differs from expected {expected_header}. Proceeding by column index matching.")

    insert_sql = """
    INSERT INTO events (id, title, category, description, event_date, event_time, venue, capacity, fee, icon)
    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """

    inserted_count = 0
    for row_data in rows[1:]:
        if not any(row_data):
            continue  # skip empty rows

        event_id = row_data[0]
        title = str(row_data[1]).strip() if row_data[1] is not None else ""
        category = str(row_data[2]).strip() if row_data[2] is not None else ""
        description = str(row_data[3]).strip() if row_data[3] is not None else ""
        event_date = str(row_data[4]).strip() if row_data[4] is not None else ""
        event_time = str(row_data[5]).strip() if row_data[5] is not None else ""
        venue = str(row_data[6]).strip() if row_data[6] is not None else ""
        capacity = int(row_data[7]) if row_data[7] is not None else 0
        fee = int(row_data[8]) if row_data[8] is not None else 0
        icon = str(row_data[9]).strip() if row_data[9] is not None else ""

        cursor.execute(insert_sql, (event_id, title, category, description, event_date, event_time, venue, capacity, fee, icon))
        inserted_count += 1

    conn.commit()
    conn.close()
    print(f"Successfully seeded {inserted_count} events from {SEED_FILE}.")

if __name__ == "__main__":
    seed_events()

    # Query and display events
    conn = get_db_connection()
    events = conn.execute("SELECT * FROM events ORDER BY id ASC").fetchall()
    print(f"\n--- Events Table in fest.db ({len(events)} rows) ---")
    for ev in events:
        print(f"[{ev['id']}] {ev['icon']} {ev['title']} ({ev['category']}) | {ev['event_date']} {ev['event_time']} @ {ev['venue']} | Cap: {ev['capacity']} | Fee: ₹{ev['fee']}")
    conn.close()
